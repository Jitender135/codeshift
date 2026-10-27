import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Sets the background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_document(md_path, docx_path):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    def flush_table(lines_to_flush):
        if not lines_to_flush:
            return
        parsed_rows = []
        for l in lines_to_flush:
            text = l.strip()
            if not text.startswith('|'):
                continue
            cells = [c.strip() for c in text.split('|')[1:-1]]
            # Skip separator line like |---|---|
            if cells and all(re.match(r'^:?-+:?$', c) for c in cells):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        col_count = max(len(r) for r in parsed_rows)
        table = doc.add_table(rows=len(parsed_rows), cols=col_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for row_idx, row in enumerate(parsed_rows):
            is_header = (row_idx == 0)
            for col_idx in range(col_count):
                cell_text = row[col_idx] if col_idx < len(row) else ""
                cell = table.cell(row_idx, col_idx)
                cell.text = ""
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

                # Shading
                if is_header:
                    set_cell_background(cell, "1F497D") # Navy blue
                elif row_idx % 2 == 1:
                    set_cell_background(cell, "F2F5F9") # Soft light blue
                else:
                    set_cell_background(cell, "FFFFFF")

                # Format text in cell
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                
                # Clean markdown links / bold in cell text
                clean_text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', cell_text)
                
                # Check formatting
                parts = re.split(r'(\*\*.*?\*\*)', clean_text)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        run = p.add_run(part)

                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)
                    else:
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def flush_code(lines_to_flush):
        if not lines_to_flush:
            return
        code_text = "".join(lines_to_flush).strip()
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F4F4F4")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith('|'):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_lines)
                table_lines = []

        # Title
        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(title_text)
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Heading 2
        elif stripped.startswith('## '):
            h2_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h2_text)
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Heading 3
        elif stripped.startswith('### '):
            h3_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h3_text)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x36, 0x5F, 0x91)

        # Heading 4
        elif stripped.startswith('#### '):
            h4_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h4_text)
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

        # Horizontal rule
        elif stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('—' * 55)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

        # Blockquote (Speaking Script / Callout)
        elif stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            # Clean quote text
            clean_q = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', quote_text)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', clean_q)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                    run.font.color.rgb = RGBColor(0x2B, 0x54, 0x7E)
                else:
                    run = p.add_run(part)
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                run.font.size = Pt(10)

        # Bullet list
        elif stripped.startswith('- ') or stripped.startswith('* '):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            clean_item = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', item_text)
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', clean_item)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
                else:
                    run = p.add_run(part)
                run.font.size = Pt(10.5)

        # Numbered list
        elif re.match(r'^\d+\.\s', stripped):
            item_text = re.sub(r'^\d+\.\s', '', stripped).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            clean_item = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', item_text)
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', clean_item)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
                else:
                    run = p.add_run(part)
                run.font.size = Pt(10.5)

        # Normal paragraph
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            clean_p = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', stripped)
            parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', clean_p)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                else:
                    run = p.add_run(part)
                run.font.size = Pt(10.5)

        i += 1

    # Final cleanup if file ends inside table or code
    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code(code_lines)

    doc.save(docx_path)
    print(f"[SUCCESS] Word document successfully saved to: {docx_path}")

if __name__ == "__main__":
    base = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base, "MIDTERM_EVALUATION_GUIDE.md")
    docx_file = os.path.join(base, "MIDTERM_EVALUATION_GUIDE.docx")
    create_styled_document(md_file, docx_file)
